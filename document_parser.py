from pathlib import Path
import csv
import io

def _decode_bytes(data):
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("latin-1", errors="replace")

def parse_uploaded_file(uploaded_file):
    name = uploaded_file.name
    ext = Path(name).suffix.lower()
    raw = uploaded_file.getvalue()
    extracted = []

    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw))
        for i, page in enumerate(reader.pages, start=1):
            txt = page.extract_text() or ""
            if txt.strip():
                extracted.append({"location": f"page {i}", "text": txt})
    elif ext == ".docx":
        from docx import Document
        doc = Document(io.BytesIO(raw))
        chunks = []
        for p in doc.paragraphs:
            if p.text.strip():
                chunks.append(p.text)
        for table in doc.tables:
            for row in table.rows:
                chunks.append(" | ".join(cell.text.strip() for cell in row.cells))
        if chunks:
            extracted.append({"location": "document text and tables", "text": "\n".join(chunks)})
    elif ext in (".txt", ".md"):
        extracted.append({"location": "text file", "text": _decode_bytes(raw)})
    elif ext == ".csv":
        text = _decode_bytes(raw)
        rows = list(csv.reader(io.StringIO(text)))
        formatted = "\n".join(" | ".join(cell for cell in row) for row in rows)
        extracted.append({"location": "CSV rows", "text": formatted})
    elif ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(values_only=True):
                vals = ["" if v is None else str(v) for v in row]
                if any(vals):
                    rows.append(" | ".join(vals))
            if rows:
                extracted.append({"location": f"sheet {ws.title}", "text": "\n".join(rows)})
    else:
        raise ValueError(f"Unsupported file type: {ext}")

    result = []
    for item in extracted:
        if item["text"].strip():
            result.append({
                "source_id": "",  # assigned consistently below
                "filename": name,
                "location": item["location"],
                "text": item["text"]
            })
    return result

def assign_source_ids(documents):
    for index, doc in enumerate(documents, start=1):
        doc["source_id"] = f"D{index:02d}"
    return documents
